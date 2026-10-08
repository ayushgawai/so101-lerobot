"""Create ACT protocol-v1.0 scoresheet: positions 1-10 x 2 trials = 20 rows.

Layout mirrors ACT_eval_scoresheet.xlsx (yellow fill cells, blue headers,
auto summary) but follows docs/eval_protocol.md instead of Fixed/Random x50.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path("ACT_positions_eval_scoresheet.xlsx")
FIRST = 19
LAST = 38


def main():
    wb = Workbook()

    header_fill = PatternFill("solid", fgColor="FF2F5496")
    header_font = Font(bold=True, color="FFFFFF")
    yellow = PatternFill("solid", fgColor="FFFFF2CC")
    gray = PatternFill("solid", fgColor="FFF2F2F2")
    title_font = Font(bold=True, size=14)
    section_font = Font(bold=True, size=11)
    thin = Border(
        left=Side(style="thin", color="FFB0B0B0"),
        right=Side(style="thin", color="FFB0B0B0"),
        top=Side(style="thin", color="FFB0B0B0"),
        bottom=Side(style="thin", color="FFB0B0B0"),
    )
    wrap = Alignment(wrap_text=True, vertical="center")

    # --- Instructions ---
    ws_i = wb.active
    ws_i.title = "Instructions"
    instructions = [
        (1, "ACT Pick-and-Place — Protocol v1.0 Evaluation Score Sheet (20 trials)", title_font),
        (3, "Aligned with docs/eval_protocol.md. Fill only the YELLOW cells on the Trials tab. Summary stats compute automatically.", None),
        (5, "BEFORE you start — preconditions:", section_font),
        (6, "1. Lock ONE checkpoint BEFORE any scored trial (record it in Trials!B3). Do not swap after seeing results.", None),
        (7, "2. Camera Pose A: capture a frame and match the Pose A reference photo; fix the mount if it does not match.", None),
        (8, "3. Cube starts on taped positions 1–10 (docs/img/init_positions.jpg). Goal = fixed bowl.", None),
        (9, "4. No manual assist. Human touch / teleop / safety stop => FAIL.", None),
        (10, "5. No time limit. End the episode (right-arrow) when success or a clear terminal failure is decided.", None),
        (12, "SUCCESS (Placed? = Y): cube is released by the policy and rests fully inside the bowl (not on rim).", section_font),
        (14, "Trial plan: 20 trials = each of positions 1–10 tested twice (trial 1 and trial 2).", None),
        (15, "Sheet order: pos1 t1, pos1 t2, pos2 t1, … pos10 t2. Reset cube to the marked tape each time.", None),
        (17, "Failure modes (pick ONE; leave blank on success):", section_font),
        (18, "  no-grasp | grasped-dropped | wrong-placement | safety-stop | other (explain in notes)", None),
        (20, "Video naming (protocol §7): videos/WXX/act_poseA_pos{NN}_t{N}.mp4", None),
        (21, "Also append each row to results/results.csv when the session is done.", None),
        (23, "How to run inference (example — confirm checkpoint lock first):", section_font),
        (24, '  .\\scripts\\run_eval.ps1 -Policy act -Mode fixed -Checkpoint "outputs/train/act_so101_pick_place_positions/checkpoints/030000/pretrained_model" -NumEpisodes 20 -RepoId "aakashv100/eval_act-pick-place-positions-poseA" -EpisodeTime -1 -ClearCache', None),
        (26, "Keyboard: right-arrow = end episode & save; Escape = stop without saving.", None),
    ]
    for row, text, font in instructions:
        ws_i[f"A{row}"] = text
        if font:
            ws_i[f"A{row}"].font = font
        else:
            ws_i[f"A{row}"].alignment = wrap
    ws_i.column_dimensions["A"].width = 140

    # --- Trials ---
    ws = wb.create_sheet("Trials")
    ws["A1"] = "ACT Evaluation — Positions 1–10 × 2 (Protocol v1.0)"
    ws["A1"].font = title_font
    ws.merge_cells("A1:I1")

    meta = [
        (3, "Checkpoint (locked)", "030000", "Train data (HF id)", "aakashv100/so101-pick-place-positions"),
        (4, "Camera pose", "A", "Domain / robot", "real / so101"),
        (5, "Operator", "", "Date (YYYY-MM-DD)", ""),
        (6, "Eval dataset repo", "aakashv100/eval_act-pick-place-positions-poseA", "Week folder (videos)", "W02"),
    ]
    for row, l1, v1, l2, v2 in meta:
        ws[f"A{row}"] = l1
        ws[f"A{row}"].font = Font(bold=True)
        ws[f"B{row}"] = v1
        ws[f"B{row}"].fill = yellow
        ws[f"D{row}"] = l2
        ws[f"D{row}"].font = Font(bold=True)
        ws[f"E{row}"] = v2
        ws[f"E{row}"].fill = yellow

    ws["A8"] = "SUMMARY (auto-computed)"
    ws["A8"].font = section_font

    grasped, placed, failure, err = "D", "E", "F", "G"
    summary = [
        (9, "Total trials", f"=COUNTA({grasped}{FIRST}:{grasped}{LAST})"),
        (10, "Successes (grasp Y AND placed Y)", f'=COUNTIFS({grasped}{FIRST}:{grasped}{LAST},"Y",{placed}{FIRST}:{placed}{LAST},"Y")'),
        (11, "Success rate", "=IF(D9=0,0,D10/D9)"),
        (12, "Grasp rate", f'=IF(D9=0,0,COUNTIF({grasped}{FIRST}:{grasped}{LAST},"Y")/D9)'),
        (13, "Mean placement err (cm, placed Y only)", f'=IFERROR(AVERAGEIFS({err}{FIRST}:{err}{LAST},{placed}{FIRST}:{placed}{LAST},"Y"),"-")'),
        (14, "Fail: no-grasp", f'=COUNTIF({failure}{FIRST}:{failure}{LAST},"no-grasp")'),
        (15, "Fail: grasped-dropped", f'=COUNTIF({failure}{FIRST}:{failure}{LAST},"grasped-dropped")'),
        (16, "Fail: wrong-placement", f'=COUNTIF({failure}{FIRST}:{failure}{LAST},"wrong-placement")'),
        (17, "Fail: safety-stop / other", f'=COUNTIF({failure}{FIRST}:{failure}{LAST},"safety-stop")+COUNTIF({failure}{FIRST}:{failure}{LAST},"other")'),
    ]
    for row, label, formula in summary:
        ws[f"A{row}"] = label
        ws.merge_cells(f"A{row}:C{row}")
        ws[f"D{row}"] = formula
        if row in (11, 12):
            ws[f"D{row}"].number_format = "0.0%"

    headers = [
        "Trial #", "init_pos_id", "trial", "Grasped? (Y/N)", "Placed? (Y/N)",
        "Failure mode", "Placement err (cm)", "video_file", "notes",
    ]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=18, column=i, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin
    ws.row_dimensions[18].height = 32

    trial_num = 1
    for pos in range(1, 11):
        for t in (1, 2):
            row = FIRST + trial_num - 1
            ws[f"A{row}"] = trial_num
            ws[f"B{row}"] = pos
            ws[f"C{row}"] = t
            for col in ("A", "B", "C"):
                ws[f"{col}{row}"].fill = gray
                ws[f"{col}{row}"].border = thin
                ws[f"{col}{row}"].alignment = Alignment(horizontal="center")
            for col in ("D", "E", "F", "G", "H", "I"):
                ws[f"{col}{row}"].fill = yellow
                ws[f"{col}{row}"].border = thin
            ws[f"H{row}"] = (
                f'=CONCATENATE("videos/",$E$6,"/act_pose",$B$4,"_pos",'
                f'TEXT(B{row},"00"),"_t",C{row},".mp4")'
            )
            trial_num += 1

    dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    dv_fail = DataValidation(
        type="list",
        formula1='"no-grasp,grasped-dropped,wrong-placement,safety-stop,other"',
        allow_blank=True,
    )
    ws.add_data_validation(dv_yn)
    ws.add_data_validation(dv_fail)
    dv_yn.add(f"D{FIRST}:E{LAST}")
    dv_fail.add(f"F{FIRST}:F{LAST}")

    ws["K8"] = "Per-position success (of 2)"
    ws["K8"].font = section_font
    for col, val in (("K9", "pos"), ("L9", "successes"), ("M9", "rate")):
        ws[col] = val
        ws[col].fill = header_fill
        ws[col].font = header_font
    for pos in range(1, 11):
        r = 9 + pos
        ws[f"K{r}"] = pos
        ws[f"L{r}"] = (
            f'=COUNTIFS($B${FIRST}:$B${LAST},K{r},$D${FIRST}:$D${LAST},"Y",'
            f'$E${FIRST}:$E${LAST},"Y")'
        )
        ws[f"M{r}"] = (
            f"=IF(COUNTIF($B${FIRST}:$B${LAST},K{r})=0,0,"
            f"L{r}/COUNTIF($B${FIRST}:$B${LAST},K{r}))"
        )
        ws[f"M{r}"].number_format = "0.0%"

    widths = {
        "A": 36, "B": 12, "C": 10, "D": 14, "E": 14, "F": 18,
        "G": 16, "H": 42, "I": 28, "K": 8, "L": 12, "M": 10,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A19"

    if OUT.exists():
        raise SystemExit(f"{OUT} already exists (delete or rename it first).")
    wb.save(OUT)
    print(f"Wrote {OUT.resolve()}")
    print(f"  Trials: {LAST - FIRST + 1} (positions 1-10 x 2)")


if __name__ == "__main__":
    main()
